"""
Script definitivo para generar el documento Word (.docx) del Anexo Metodológico CarBot,
estructurado limpiamente en 6 secciones principales sin relleno teórico y 100% técnico.
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    cell._tc.get_or_add_tcPr().append(tc_mar)

def set_table_borders(table, color="D3D3D3"):
    tbl_pr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="1F4E79"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="1F4E79"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tbl_pr.append(borders)

def add_callout_box(doc, text_paragraphs, title=None, fill_hex="F2F5F9", border_color="1F4E79"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, fill_hex)
    set_cell_margins(cell, top=130, bottom=130, left=170, right=130)
    
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    if title:
        run_title = p.add_run(f"{title}\n")
        run_title.bold = True
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(10)
        run_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        
    for i, line in enumerate(text_paragraphs):
        if i == 0 and not title:
            run = p.add_run(line)
        else:
            p_sub = cell.add_paragraph()
            p_sub.paragraph_format.space_before = Pt(1)
            p_sub.paragraph_format.space_after = Pt(1)
            p_sub.paragraph_format.line_spacing = 1.15
            run = p_sub.add_run(line)
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_heading(p, text, level=1):
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.bold = True
    if level == 1:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    elif level == 2:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
    return r

def build_table(doc, headers, data, col_widths):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    for col_idx, header_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.width = col_widths[col_idx]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(header_text)
        r.font.name = "Calibri"
        r.font.size = Pt(9)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for row_idx, row_data in enumerate(data):
        bg = "FFFFFF" if row_idx % 2 == 0 else "F7F9FC"
        for col_idx, text_val in enumerate(row_data):
            cell = table.cell(row_idx + 1, col_idx)
            cell.width = col_widths[col_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=90, right=90)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(text_val)
            r.font.name = "Calibri"
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(0x26, 0x26, 0x26)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table

def build_docx(output_path):
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        hp = section.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("ANEXO METODOLÓGICO: SCRUM + CRISP-DM — CARBOT 2026")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(0x8A, 0x9B, 0xA8)
        
        fp = section.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Facultad de Ingeniería de Sistemas e Informática — Tesis Profesional")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(0x8A, 0x9B, 0xA8)

    # TÍTULO PRINCIPAL
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_pre = title_p.add_run("ANEXO METODOLÓGICO\n")
    run_pre.font.name = "Calibri"
    run_pre.font.size = Pt(12)
    run_pre.font.bold = True
    run_pre.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

    run_title = title_p.add_run("METODOLOGÍA DE DESARROLLO DEL SISTEMA CARBOT MEDIANTE LA INTEGRACIÓN HÍBRIDA SCRUM Y CRISP-DM")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(17)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    meta_box = [
        "Título de la Tesis: \"Chatbot utilizando machine learning para el diagnóstico vehicular en talleres mecánicos en Carabayllo 2026\"",
        "Sede de Aplicación Experimental: Taller Automotriz CARTER MOTOR'S E.I.R.L. (Carabayllo, Lima - Perú)",
        "Diseño de la Investigación: Preexperimental con Pre-test (O₁) y Post-test (O₂) sobre muestra de 60 casos reales",
        "Clasificación del Documento: Anexo Metodológico y Técnico de Ingeniería de Software e Inteligencia Artificial"
    ]
    add_callout_box(doc, meta_box, title="FICHA TÉCNICA DE LA INVESTIGACIÓN", fill_hex="EBF2FA", border_color="1F4E79")

    # 1. QUÉ RESUELVE (Exactamente 1 párrafo de 7-8 líneas)
    format_heading(doc.add_paragraph(), "1. Qué resuelve", level=1)
    p1 = doc.add_paragraph()
    p1.paragraph_format.line_spacing = 1.15
    p1.paragraph_format.space_after = Pt(8)
    p1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r1 = p1.add_run(
        "En el taller mecánico CARTER MOTOR'S E.I.R.L. de Carabayllo, el diagnóstico vehicular se realizaba tradicionalmente de forma "
        "empírica y manual, lo que generaba demoras promedio superiores a 40 minutos en inspección preliminar, dispersión de criterios "
        "entre técnicos y ausencia de registros estandarizados de fallas. Para solucionar esta problemática, el proyecto desarrolló "
        "CarBot, una plataforma integral que asiste al mecánico en tiempo real vía WhatsApp y centraliza el registro experimental de la tesis. "
        "El sistema implementa procesamiento de lenguaje natural con traducción de jerga automotriz, clasificación multiclase mediante "
        "Linear SVM con vectorización TF-IDF jerárquica en 48 averías, enriquecimiento procedimental con RAG (FAISS) y un panel web en "
        "React/FastAPI con PostgreSQL para la medición preexperimental comparativa (Pre-test y Post-test) sobre una muestra de 60 casos reales."
    )
    r1.font.name = "Calibri"
    r1.font.size = Pt(10)

    # 2. ALCANCE Y ARQUITECTURA DEL SISTEMA
    format_heading(doc.add_paragraph(), "2. Alcance y Arquitectura del Sistema", level=1)
    
    # 2.1 Módulos y Tecnologías (Tabla Unificada)
    format_heading(doc.add_paragraph(), "2.1 Módulos y Tecnologías del Sistema", level=2)
    modulos_headers = ["Módulo", "Qué hace", "Tecnología / Componente Clave"]
    modulos_data = [
        ("Asistencia vía WhatsApp", "Recepción de consultas por WhatsApp Cloud API, acuse preliminar inmediato y entrega de reporte con Top 3 fallas y pruebas físicas.", "FastAPI Webhooks, Meta Cloud API, Worker asíncrono."),
        ("Procesamiento NLP", "Traducción de jerga peruana ('cascabelea', 'se chupa'), extracción de códigos DTC OBD-II (P0301, P0420) y normalización léxica.", "traductor_jerga.py, semantic_purifier.py, text_processor.py."),
        ("Clasificación ML", "Vectorización TF-IDF y clasificación multiclase con Linear SVM jerárquico calibrado (7 Macro-Sistemas y 48 averías específicas).", "linear_svm_diagnostico.joblib, tfidf_diagnostico.joblib, Scikit-Learn."),
        ("Base RAG (FAISS)", "Recuperación contextual de tolerancias eléctricas (13.8V - 14.4V), torques y pruebas metrológicas desde manuales de servicio OEM.", "index.faiss, metadata.json, FAISS Vectorstore."),
        ("Instrumentos de Validación (Anexo 2)", "Registro y control de fichas Pre-test (manual) y Post-test (asistido por CarBot) evaluando PPCF (Ficha 1), PRDC (Ficha 2) y TPRD (Ficha 3).", "React 18, TypeScript, SQLAlchemy async, validaciones_taller."),
        ("Aislamiento y Exportación", "Blindaje contra contaminación entre datos PILOTO y OFICIAL (0/60) y exportación oficial en formato LONG (CSV) para contraste estadístico.", "PostgreSQL 15 (carbot_db), Alembic 20260919_04, Pandas/CSV.")
    ]
    build_table(doc, modulos_headers, modulos_data, [Inches(1.8), Inches(3.0), Inches(1.7)])

    # 2.2 Fuera de alcance
    format_heading(doc.add_paragraph(), "2.2 Fuera de alcance", level=2)
    fuera_items = [
        "• Interacción directa con clientes o conductores particulares: Diseñado para uso exclusivo del personal técnico en bahía de trabajo.",
        "• Sustitución de pruebas físicas o metrológicas: CarBot sugiere hipótesis y pruebas metrológicas, pero la confirmación física real siempre la realiza el mecánico en taller.",
        "• Reentrenamiento automático desatendido: El modelo Linear SVM y el índice FAISS están congelados criptográficamente (CARBOT_PRECAMPO_FINAL_HASH_MANIFEST.json) para garantizar reproducibilidad en la tesis.",
        "• Alteración de registros históricos: La base de datos no realiza borrados físicos de fichas auditadas, preservando la trazabilidad."
    ]
    for item in fuera_items:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.line_spacing = 1.15
        p_item.paragraph_format.space_after = Pt(2)
        r = p_item.add_run(item)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)

    # 2.3 Roles
    format_heading(doc.add_paragraph(), "2.3 Roles del Sistema y del Equipo", level=2)
    usuarios_headers = ["Rol en Plataforma / Equipo", "Responsabilidad Principal"]
    usuarios_data = [
        ("Mecánico de Taller", "Consulta diagnósticos en tiempo real por WhatsApp, reporta síntomas coloquiales o códigos DTC y ejecuta las pruebas de confirmación física en el vehículo."),
        ("Investigador / Tesista (Product Owner)", "Administra el panel web de validación, registra las fichas de Pre-test y Post-test, audita la completitud de campos, controla tiempos y exporta datos oficiales."),
        ("Asesor de Tesis / Administrador (Scrum Master)", "Supervisa el avance del trabajo de campo (0/60 casos), audita la integridad de datos, verifica el aislamiento del entorno piloto y valida las métricas de investigación.")
    ]
    build_table(doc, usuarios_headers, usuarios_data, [Inches(2.2), Inches(4.3)])

    # 2.4 Diagramas Principales
    format_heading(doc.add_paragraph(), "2.4 Diagramas Principales del Sistema", level=2)
    flujo_box = [
        "1. Solicitud de Entrada: El mecánico envía un mensaje de texto con síntomas coloquiales o código DTC vía WhatsApp.",
        "2. Acuse Inmediato: FastAPI Webhook recibe el evento HTTPS y responde HTTP 200 a Meta en < 2 segundos ('⏳ Analizando consulta...').",
        "3. Pipeline NLP: Traductor de jerga normaliza modismos peruanos y purificador semántico aísla códigos DTC OBD-II.",
        "4. Inferencia Jerárquica: Linear SVM clasifica el Macro-Sistema (Nivel 1) y calcula probabilidades calibradas de 48 averías (Nivel 2).",
        "5. Evaluación de Confianza: Si margen Δ ≥ 10% o existe DTC, pasa directo a reporte. Si es ambiguo (Δ < 10%), genera repregunta técnica.",
        "6. Enriquecimiento RAG: Motor FAISS recupera tolerancias OEM, torques y pruebas de descarte metrológico.",
        "7. Entrega y Registro: Envío del reporte por WhatsApp y persistencia en tabla diagnosticos para vincularlo a la Ficha Post-test."
    ]
    add_callout_box(doc, flujo_box, title="FIGURA 1. FLUJO DEL PROCESAMIENTO DIAGNÓSTICO EN CARBOT", fill_hex="F7F9FC", border_color="1F4E79")

    arq_box = [
        "• Capa de Clientes: WhatsApp en teléfonos móviles de mecánicos + Panel Web en React/Vite/TypeScript para el investigador.",
        "• Comunicaciones Seguras: Meta WhatsApp Cloud API mediante webhooks cifrados HTTPS.",
        "• Capa de Lógica y Servicios (FastAPI / Python 3.11): API REST modular, worker asíncrono de colas, motor NLP y controlador diagnóstico.",
        "• Capa de Machine Learning & RAG: Clasificador Linear SVM (.joblib), vectorizador TF-IDF (.joblib) e índice vectorial FAISS (.faiss).",
        "• Capa de Persistencia: Base de datos relacional PostgreSQL 15 (carbot_db) con migraciones versionadas en Alembic (20260919_04)."
    ]
    add_callout_box(doc, arq_box, title="FIGURA 2. ARQUITECTURA DESPLEGADA DE CARBOT (FASTAPI + REACT + POSTGRESQL)", fill_hex="F7F9FC", border_color="1F4E79")

    # 3. MODELADO PREDICTIVO Y REGLAS CANÓNICAS (CRISP-DM)
    format_heading(doc.add_paragraph(), "3. Modelado Predictivo y Reglas Canónicas de Decisión (CRISP-DM)", level=1)
    p_ml_desc = doc.add_paragraph()
    p_ml_desc.paragraph_format.line_spacing = 1.15
    p_ml_desc.paragraph_format.space_after = Pt(4)
    p_ml_desc.add_run(
        "El modelado predictivo se ejecutó sobre un dataset canónico de 5,374 casos estructurados con 7 Macro-Sistemas y 48 averías específicas. "
        "El clasificador seleccionado es Linear SVM con probabilidades calibradas mediante Platt Scaling (CalibratedClassifierCV). "
        "Las decisiones diagnósticas se rigen por tres reglas matemáticas canónicas:"
    )

    add_callout_box(
        doc,
        [
            "P_comb(Fᵢ) = P(Fᵢ) × [P(S_k)]^γ",
            "Donde γ = 0.65 es el exponente calibrado empíricamente que penaliza hipótesis fuera de sistema sin castigar averías que comparten sintomatología entre subsistemas conexos."
        ],
        title="ECUACIÓN 1: PONDERACIÓN JERÁRQUICA SUAVE DE CARBOT",
        fill_hex="F7F9FC",
        border_color="2E75B6"
    )

    add_callout_box(
        doc,
        [
            "P_dtc(Fᵢ) = min(1.0, P_comb(Fᵢ) × 2.5)   para toda falla Fᵢ candidata del DTC.",
            "Garantiza que la evidencia electrónica digital del escáner OBD-II prevalezca sobre la subjetividad coloquial.",
            "\nRegla de Margen Canónico: Δ = P(Top 1) - P(Top 2). Si Δ ≥ 10% o P(Top 1) ≥ 70%, responde directamente; si no, activa repregunta."
        ],
        title="ECUACIÓN 2: MULTIPLICADOR DE AUTORIDAD DTC Y MARGEN CANÓNICO",
        fill_hex="F7F9FC",
        border_color="2E75B6"
    )

    p_eval = doc.add_paragraph()
    p_eval.paragraph_format.line_spacing = 1.15
    p_eval.paragraph_format.space_after = Pt(4)
    p_eval.add_run(
        "Resultados en Holdout Ciego (20%): Exactitud en 48 Averías: 93.07% | Exactitud en Macro-Sistema: 98.71% | "
        "Top-3 Accuracy: 98.71% | F1-Score Macro: > 0.91 | Congelamiento Criptográfico: 13 firmas SHA-256 inmutables "
        "(CARBOT_PRECAMPO_FINAL_HASH_MANIFEST.json)."
    )

    # 4. MATRIZ CRISP-DM ⟷ SCRUM Y DEFINITION OF DONE (DoD)
    format_heading(doc.add_paragraph(), "4. Matriz Metodológica (CRISP-DM ⟷ SCRUM) y Definition of Done (DoD)", level=1)
    
    format_heading(doc.add_paragraph(), "4.1 Cronograma de Integración por Sprints", level=2)
    sprints_headers = ["Sprint", "Duración", "Fases CRISP-DM", "Objetivo del Sprint (Sprint Goal)", "Entregables e Incrementos de Software"]
    sprints_data = [
        ("Sprint 0", "2 sem.", "I. Comprensión Negocio", "Levantamiento de requisitos de taller, arquitectura base y variables de tesis.", "• Documento de Arquitectura.\n• Matriz de variables e hipótesis.\n• Entorno base FastAPI + Postgres + React."),
        ("Sprint 1", "2 sem.", "II. Comprensión Datos", "Recolección del corpus automotriz y diseño de la taxonomía jerárquica de 48 averías.", "• Dataset canónico preliminar (5,374 casos).\n• Taxonomía 7 Macro-Sistemas y 48 averías.\n• Diccionario de jerga automotriz."),
        ("Sprint 2", "2 sem.", "III. Preparación Datos", "Construcción del pipeline de NLP (traductor, purificador, procesador) y TF-IDF.", "• text_processor.py y traductor_jerga.py.\n• Extractor de códigos DTC.\n• tfidf_diagnostico.joblib (1, 2 n-gramas)."),
        ("Sprint 3", "2 sem.", "IV. Modelado Predictivo", "Entrenamiento del clasificador Linear SVM multiclase, calibración Platt y Macro-Sistema.", "• linear_svm_diagnostico.joblib.\n• Modelo de Macro-Sistemas.\n• Fusión jerárquica con γ = 0.65."),
        ("Sprint 4", "2 sem.", "IV (RAG) y V (Evaluación)", "Construcción de vectorstore FAISS con manuales OEM, fusión DTC y holdout.", "• Índice index.faiss y metadata.json.\n• politica_fusion.py (DTC 2.5×).\n• Holdout verificado: 93.07% Accuracy."),
        ("Sprint 5", "2 sem.", "VI. Despliegue Backend", "API REST en FastAPI, PostgreSQL relacional, migraciones Alembic y WhatsApp.", "• Endpoints /diagnostico y /webhook.\n• Worker de colas y rate limiter.\n• Conexión interactiva validada con Meta API."),
        ("Sprint 6", "2 sem.", "VI. Despliegue Frontend", "Panel web React/Vite para visualización y fichas técnicas de taller (Anexo 2).", "• Interfaz responsive en React/TypeScript.\n• Módulo de Fichas 1 (PPCF), 2 (PRDC) y 3 (TPRD).\n• Medición de completitud y tiempos."),
        ("Sprint 7", "2 sem.", "V y VI: Hardening Final", "Congelamiento criptográfico, aislamiento de muestra (0/60) y check pre-campo.", "• Manifiesto SHA-256 inmutable.\n• Aislamiento de casos PILOTO vs OFICIAL.\n• Exportador LONG CSV de Anexo 2.\n• Dictamen APTO PARA PILOTO.")
    ]
    build_table(doc, sprints_headers, sprints_data, [Inches(0.8), Inches(0.6), Inches(1.3), Inches(1.9), Inches(1.9)])

    format_heading(doc.add_paragraph(), "4.2 Criterios Técnicos de Definition of Done (DoD)", level=2)
    dod_box = [
        "1. Modularidad Estricta: Ningún archivo supera las 400 líneas de código (Clean Architecture).",
        "2. Pruebas Unitarias Backend: Aprobadas al 100% en Pytest (pytest -q).",
        "3. Cobertura de Código: Cobertura superior al 55% en backend/src.",
        "4. Frontend Verificado: Tipado TypeScript sin errores (npm run build) y arnés React aprobado (npm run test:harness).",
        "5. Cero Regresiones: Ausencia total de regresiones en funcionalidades existentes.",
        "6. Esquema de Base de Datos: Migraciones Alembic al día en head (versión 20260919_04).",
        "7. Congelamiento ML: Inmutabilidad estricta de los modelos contra el manifiesto SHA-256."
    ]
    add_callout_box(doc, dod_box, title="CRITERIOS TÉCNICOS DE DEFINITION OF DONE (DoD)", fill_hex="F2F5F9", border_color="1F4E79")

    # 5. HISTORIAS DE USUARIO (Alineadas a la pantalla real de tesis)
    format_heading(doc.add_paragraph(), "5. Historias de Usuario (Alineadas a los Instrumentos de Tesis)", level=1)
    hu_headers = ["ID", "Historia de Usuario", "Módulo / Ficha Vinculada", "Criterios de Aceptación Técnicos"]
    hu_data = [
        ("HU-01", "Como mecánico, quiero redactar síntomas con jerga peruana ('cascabelea', 'se chupa'), para que el sistema reconozca la falla sin tecnicismos.", "NLP y Jerga", "1. Mapeo a términos estandarizados.\n2. Eliminación de tildes y signos.\n3. Latencia léxica < 5 ms."),
        ("HU-02", "Como mecánico, quiero ingresar un código OBD-II (ej. P0301) junto al síntoma, para que el sistema priorice la evidencia digital del escáner.", "Purificador Semántico", "1. Aislamiento exacto del patrón DTC.\n2. Multiplicador de autoridad 2.5×.\n3. Fijación del Macro-Sistema."),
        ("HU-03", "Como técnico en bahía, quiero consultar el diagnóstico desde mi celular por WhatsApp, para obtener orientación rápida sin ir a una PC.", "Asistencia WhatsApp", "1. Acuse preliminar en < 2 segundos.\n2. Reporte con Top 3 fallas y pruebas.\n3. Fallback degradado local si la API demora."),
        ("HU-04", "Como mecánico, quiero recibir tolerancias de fábrica (13.8V - 14.4V) y pruebas OEM, para descartar la falla con instrumental de taller.", "Base RAG (FAISS)", "1. Recuperación de manuales OEM con FAISS.\n2. Inclusión de rangos de carga y torques.\n3. Sugerencia de prueba física metrológica."),
        ("HU-05", "Como investigador, quiero registrar fichas Pre-test del método tradicional sin CarBot, para capturar la línea base del taller sin sesgo.", "Validación Pre-test", "1. Formulario autónomo sin diagnostico_id.\n2. Registro de hipótesis y falla real confirmada.\n3. Cronometraje manual en minutos."),
        ("HU-06", "Como investigador, quiero vincular una consulta de WhatsApp a la ficha Post-test, para auditar la Ficha 1 (PPCF) con inmutabilidad de predicción.", "Ficha 1: Predicción (PPCF)", "1. Precarga automática de predicción y confianza.\n2. Falla predicha bloqueada en backend.\n3. Asociación atómica del diagnostico_id."),
        ("HU-07", "Como investigador, quiero auditar los campos obligatorios del registro técnico, para medir el ratio de completitud de la Ficha 2 (PRDC).", "Ficha 2: Información (PRDC)", "1. Validación de 8 campos estructurados mínimos.\n2. Cálculo automático del indicador PRDC.\n3. Almacenamiento binario del cumplimiento."),
        ("HU-08", "Como investigador, quiero registrar el tiempo del proceso diagnóstico, para calcular el tiempo promedio de respuesta de la Ficha 3 (TPRD).", "Ficha 3: Tiempo (TPRD)", "1. Registro independiente de tiempo_diagnostico_minutos.\n2. Separación de telemetría y latencia ML.\n3. Cálculo de la sumatoria y promedio."),
        ("HU-09", "Como tesista, quiero disponer de un entorno PILOTO independiente del OFICIAL, para que los ensayos de práctica no contaminen la muestra de 60 casos.", "Aislamiento de Datos", "1. Selector explícito (PILOTO vs OFICIAL).\n2. Modal de confirmación obligatorio para oficial.\n3. Muestra oficial en 0/60 previa al campo."),
        ("HU-10", "Como asesor o tesista, quiero exportar la data en formato LONG cumpliendo el Anexo 2, para realizar el análisis estadístico inferencial.", "Exportador Oficial", "1. Exportación CSV con columnas metodológicas.\n2. Inclusión exclusiva de casos verificados oficiales.\n3. Exclusión total de desarrollo y piloto.")
    ]
    build_table(doc, hu_headers, hu_data, [Inches(0.6), Inches(2.4), Inches(1.4), Inches(2.1)])

    # 6. GESTIÓN DE RIESGOS E IMPEDIMENTOS TÉCNICOS
    format_heading(doc.add_paragraph(), "6. Gestión de Riesgos e Impedimentos Técnicos", level=1)
    riesgos_headers = ["Riesgo / Impedimento Detectado", "Impacto en la Investigación", "Estrategia de Mitigación en SCRUM"]
    riesgos_data = [
        ("Fuga de Información (Data Leakage) en ML", "Sobreestimación artificial del rendimiento del modelo en producción.", "Ajuste (fit) de TF-IDF exclusivo en 80% train. Validación ciega sobre el 20% holdout sin contacto previo."),
        ("Caídas o cuotas en APIs externas de lenguaje", "Bloqueo del mecánico en WhatsApp y consultas huérfanas sin diagnóstico.", "Invariante de entrega con fallback degradado local: fusión directa de Linear SVM y FAISS en milisegundos."),
        ("Contaminación de Muestra Oficial (60 casos)", "Invalidez académica de los resultados preexperimentales de la tesis.", "Triple barrera: Modal en Frontend, validación de tipo_registro en Backend y CheckConstraint en PostgreSQL."),
        ("Manipulación accidental de modelos congelados", "Pérdida de trazabilidad y reproducibilidad ante el jurado de tesis.", "Manifiesto criptográfico de 13 artefactos sellados con firmas SHA-256 (CARBOT_PRECAMPO_FINAL_HASH_MANIFEST.json)."),
        ("Confusión entre latencia computacional y tiempo de taller", "Distorsión metodológica en la Ficha 3 (eficiencia del diagnóstico).", "Separación estricta de columnas: tiempo_inferencia_ml_ms, duracion_sistema_segundos y tiempo_diagnostico_minutos.")
    ]
    build_table(doc, riesgos_headers, riesgos_data, [Inches(1.8), Inches(2.0), Inches(2.7)])

    # Guardado con manejo de bloqueo
    try:
        doc.save(output_path)
        print(f"Documento Word limpio generado con éxito en: {output_path}")
    except PermissionError:
        alt_path = output_path.replace(".docx", "_SIN_RELLENO.docx")
        doc.save(alt_path)
        print(f"AVISO: El archivo destino estaba abierto en Word. Se guardó exitosamente en: {alt_path}")
        return alt_path
    return output_path

if __name__ == "__main__":
    target = os.path.abspath("docs/anexos/ANEXO_METODOLOGIA_DESARROLLO_SCRUM_CRISP_DM.docx")
    build_docx(target)
